"""按用户截图生成文字和列表均可编辑的 Word 简历。"""
import json
from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
BLUE = '294E70'
GRAY = '606060'
BODY_SIZE = 11.5
BODY_LINE = 16.9


def element(tag, **attributes):
    """创建需要显式设置的 Word XML 节点。"""
    node = OxmlElement(tag)
    for name, value in attributes.items():
        node.set(qn('w:' + name), str(value))
    return node


def format_run(run, size=BODY_SIZE, bold=False, color=GRAY, symbol=False):
    """同时设置中文和西文字体，使 Word 保持一致的文字显示。"""
    run.font.name = 'Segoe UI Symbol' if symbol else 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), 'Segoe UI Symbol' if symbol else 'Arial')
    fonts.set(qn('w:hAnsi'), 'Segoe UI Symbol' if symbol else 'Arial')
    fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    return run


def paragraph_format(paragraph, line=BODY_LINE, before=0, after=0, keep=False):
    """控制自然段的间距，保留可正常换行的正文。"""
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line)
    fmt.keep_with_next = keep
    fmt.widow_control = True
    paragraph._p.get_or_add_pPr().append(element('w:contextualSpacing', val='0'))


def cell_properties(cell, fill=None, left=0, right=0, top=0, bottom=0):
    """生成参考图使用的无边框布局，标题格按需填充深蓝。"""
    props = cell._tc.get_or_add_tcPr()
    margins = element('w:tcMar')
    for side, value in [('left', left), ('right', right), ('top', top), ('bottom', bottom)]:
        margins.append(element('w:' + side, w=value, type='dxa'))
    props.append(margins)
    borders = element('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        borders.append(element('w:' + side, val='nil'))
    props.append(borders)
    if fill:
        props.append(element('w:shd', val='clear', fill=fill))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table_layout(table, widths):
    """固定列宽，让照片、个人信息和标题条位置与参考图接近。"""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    props = table._tbl.tblPr
    props.append(element('w:tblInd', w=0, type='dxa'))
    borders = element('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        borders.append(element('w:' + side, val='nil'))
    props.append(borders)
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
        for row in table.rows:
            row.cells[index].width = Inches(width)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(element('w:cantSplit'))


def spacer(doc, points):
    """添加用于参考版式的空白，不把段落推到另一页。"""
    p = doc.add_paragraph()
    paragraph_format(p, line=points)
    format_run(p.add_run(''), size=1)
    return p


def bullet_numbering(doc):
    """建立 Word 原生项目符号，文字编辑后续行仍正确缩进。"""
    numbering = doc.part.numbering_part.element
    abstract = element('w:abstractNum', abstractNumId=90)
    abstract.append(element('w:multiLevelType', val='singleLevel'))
    level = element('w:lvl', ilvl=0)
    level.append(element('w:start', val=1))
    level.append(element('w:numFmt', val='bullet'))
    level.append(element('w:lvlText', val='•'))
    level.append(element('w:lvlJc', val='left'))
    ppr = element('w:pPr')
    ppr.append(element('w:ind', left=300, hanging=210))
    tabs = element('w:tabs')
    tabs.append(element('w:tab', val='num', pos=300))
    ppr.append(tabs)
    level.append(ppr)
    rpr = element('w:rPr')
    rpr.append(element('w:rFonts', ascii='Arial', hAnsi='Arial'))
    rpr.append(element('w:color', val=GRAY))
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)
    num = element('w:num', numId=90)
    num.append(element('w:abstractNumId', val=90))
    numbering.append(num)
    return 90


def add_bullet(doc, text, number_id, bold_label=False):
    """写入原文，项目符号只用于版式而不变更句子内容。"""
    p = doc.add_paragraph()
    paragraph_format(p)
    num = element('w:numPr')
    num.append(element('w:ilvl', val=0))
    num.append(element('w:numId', val=number_id))
    p._p.get_or_add_pPr().append(num)
    if bold_label:
        label, rest = text.split('：', 1)
        format_run(p.add_run(label + '：'), bold=True)
        format_run(p.add_run(rest))
    else:
        format_run(p.add_run(text))


def add_header(doc, content, photo_path):
    """创建照片、姓名和双列个人信息，全部文字均为可编辑内容。"""
    table = doc.add_table(rows=1, cols=2)
    table_layout(table, [1.57, 6.34])
    photo_cell, info_cell = table.rows[0].cells
    cell_properties(photo_cell)
    cell_properties(info_cell)
    p = photo_cell.paragraphs[0]
    paragraph_format(p, line=BODY_LINE)
    # 图片所在段落使用最小行高，让 Word 按照片高度展开，避免固定行高裁掉人像。
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.left_indent = Inches(.21)
    p.add_run().add_picture(str(photo_path), width=Inches(.86), height=Inches(1.08))
    inline = p.runs[-1]._r.xpath('.//wp:docPr')[0]
    inline.set('descr', '个人照片，可在 Word 或 WPS 中更换图片')
    title = info_cell.paragraphs[0]
    title.style = doc.styles['Normal']
    paragraph_format(title, line=32, after=10)
    format_run(title.add_run(content['name']), size=20, bold=True, color=BLUE)
    detail = info_cell.add_table(rows=2, cols=2)
    table_layout(detail, [3.15, 3.19])
    records = [
        ('▦', '年　龄： ' + content['age']), ('☎', '电　话： ' + content['phone']),
        ('✉', '邮　箱： ' + content['email']), ('★', '本　硕： ' + content['education'])
    ]
    for cell, (icon, text) in zip([detail.cell(0, 0), detail.cell(0, 1), detail.cell(1, 0), detail.cell(1, 1)], records):
        cell_properties(cell)
        p = cell.paragraphs[0]
        paragraph_format(p, line=18)
        format_run(p.add_run(icon + '　'), size=11, color=BLUE, symbol=True)
        format_run(p.add_run(text), size=10.3, color='404040')
    paragraph_format(info_cell.paragraphs[-1], line=1)
    format_run(info_cell.paragraphs[-1].add_run(''), size=1)


def add_project_heading(doc, title, width):
    """按参考图创建蓝色标签及右侧细灰线，标题仍是 Word 原生文本。"""
    table = doc.add_table(rows=1, cols=2)
    table_layout(table, [width, 7.91 - width])
    label, line = table.rows[0].cells
    cell_properties(label, fill=BLUE, left=60)
    cell_properties(line)
    p = label.paragraphs[0]
    paragraph_format(p, line=21, keep=True)
    format_run(p.add_run('› '), size=16, bold=True, color='FFFFFF')
    format_run(p.add_run(title), size=11.3, bold=True, color='FFFFFF')
    p = line.paragraphs[0]
    paragraph_format(p, line=1, keep=True)
    borders = element('w:pBdr')
    borders.append(element('w:bottom', val='single', sz=4, space=0, color='A0A0A0'))
    p._p.get_or_add_pPr().append(borders)
    format_run(p.add_run(''), size=1)


def main():
    """提取参考截图中的照片，并保存一份完整可编辑简历。"""
    content = json.loads((WORK / 'resume-content.json').read_text(encoding='utf-8'))
    photo_path = WORK / 'photo-from-reference.png'
    screenshot = Image.open('C:/Users/19344/AppData/Local/Temp/codex-clipboard-5eb75546-fd96-4dc9-8f9b-8d0c56904b17.png')
    screenshot.crop((86, 54, 213, 213)).save(photo_path)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.32)
    section.bottom_margin = Inches(.35)
    section.left_margin = Inches(.36)
    section.right_margin = Inches(.23)
    section.header_distance = Inches(.1)
    section.footer_distance = Inches(.1)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(GRAY)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(BODY_LINE)
    add_header(doc, content, photo_path)
    spacer(doc, 26)
    numbering = bullet_numbering(doc)
    for index, project in enumerate(content['projects']):
        if index:
            spacer(doc, 15)
        add_project_heading(doc, project['title'], 2.15 if index == 0 else 3.82)
        p = doc.add_paragraph()
        paragraph_format(p, before=8)
        format_run(p.add_run(project['description']))
        for bullet_index, text in enumerate(project['bullets']):
            add_bullet(doc, text, numbering, bold_label=index == 1 and bullet_index == 0)
    doc.core_properties.title = '韩克迪个人简历'
    doc.core_properties.subject = '大模型工程知识学习助手与电商数据分析 Agent'
    doc.core_properties.author = content['name']
    doc.core_properties.last_modified_by = content['name']
    output = ROOT / 'output/简历/韩克迪_简历.docx'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == '__main__':
    main()
